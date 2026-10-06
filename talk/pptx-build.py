# Builds the COVESA-template PowerPoint from talk/out/pptx (see talk/out/pptx-capture.mjs).
# Usage: python talk/pptx-build.py <template.pptx> <out.pptx>
import copy
import json
import sys
from pathlib import Path

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

src = Path(__file__).parent / 'out' / 'pptx'
slides = json.loads((src / 'slides.json').read_text())
prs = Presentation(sys.argv[1])
content_layout, section_layout = prs.slide_masters[1].slide_layouts[0], prs.slide_masters[1].slide_layouts[1]

AREA_L, AREA_T, AREA_W, AREA_H = Inches(0.67), Inches(1.25), Inches(12.0), Inches(5.6)


def set_text(shape, text):
	tf = shape.text_frame
	for p in tf.paragraphs[1:]:
		p._p.getparent().remove(p._p)
	runs = tf.paragraphs[0].runs
	for r in runs[1:]:
		r._r.getparent().remove(r._r)
	first, *rest = text.split('\n')
	runs[0].text = first
	prev = runs[0]._r
	for line in rest:
		br = prev.makeelement(qn('a:br'), {})
		prev.addnext(br)
		prev = copy.deepcopy(runs[0]._r)
		prev.find(qn('a:t')).text = line
		br.addnext(prev)


TITLE_FONT = ImageFont.truetype('/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts/arialbd.ttf', 100)
TITLE_WIDTH_PT = (10752274 - 2 * 91425) / 12700


def fit_title(text):
	"""Largest size up to 36 pt that keeps the title on one line, or None if that needs under 24 pt.
	Measured with Arial Bold plus 10%, since the template's Open Sans runs wider."""
	width_at_1pt = TITLE_FONT.getlength(text) / 100 * 1.1
	size = min(36, int(TITLE_WIDTH_PT / width_at_1pt))
	return size if size >= 24 else None


def add_notes(slide, notes):
	if notes:
		slide.notes_slide.notes_text_frame.text = notes


title_slide, thanks_slide = prs.slides[0], prs.slides[6]
first = slides[0]
title_text = {
	'TITLE': ('How GM uses vsomeip\nin its ADAS stack', 1.9, 1.25, None),
	'Sub-head': ('Copies, discovery, ownership and threads: proposals for upstream', 3.25, 0.75, 18),
	'Date': ('Siddharth Menon, General Motors  ·  28-29 October 2026', 4.1, 0.45, 16),
}
for sh in title_slide.shapes:
	if sh.has_text_frame and sh.text_frame.text in title_text:
		text, top, height, size = title_text[sh.text_frame.text]
		set_text(sh, text)
		sh.top, sh.height = Inches(top), Inches(height)
		if size:
			for r in sh.text_frame.paragraphs[0].runs:
				r.font.size = Pt(size)
add_notes(title_slide, first['notes'])

for s in slides[1:]:
	if s['type'] == 'divider':
		slide = prs.slides.add_slide(section_layout)
		slide.shapes.title.text = s['title']
		sub = slide.placeholders[1]
		sub.text = s['subtitle']
		if '\n' in s['subtitle']:
			lift = Inches(1.0)
			for ph in (slide.shapes.title, sub):
				ph.left, ph.top, ph.width, ph.height = ph.left, ph.top - lift, ph.width, ph.height
			sub.height += lift
			for i, p in enumerate(sub.text_frame.paragraphs):
				for r in p.runs:
					r.font.size = Pt(18 if i == 0 else 13)
	else:
		slide = prs.slides.add_slide(content_layout)
		title = slide.shapes.title
		title.text = s['title']
		size, kicker_top = fit_title(s['title']), Inches(0.12)
		if size is None:
			size, kicker_top = 26, Inches(0.04)
			title.left, title.top, title.width, title.height = Emu(608086), Inches(0.38), Emu(10752274), Inches(0.75)
		title.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
		body = slide.placeholders[1]
		body._element.getparent().remove(body._element)
		if s['kicker']:
			tb = slide.shapes.add_textbox(AREA_L, kicker_top, AREA_W, Inches(0.3))
			r = tb.text_frame.paragraphs[0].add_run()
			r.text = s['kicker'].upper()
			r.font.size, r.font.bold, r.font.color.rgb = Pt(11), True, RGBColor(0x00, 0x95, 0xA6)
		w, h = Image.open(src / s['image']).size
		scale = min(AREA_W / w, AREA_H / h)
		pw, ph = int(w * scale), int(h * scale)
		slide.shapes.add_picture(str(src / s['image']), AREA_L + (AREA_W - pw) // 2, AREA_T, Emu(pw), Emu(ph))
	add_notes(slide, s['notes'])

ids = prs.slides._sldIdLst
for i in (5, 4, 3, 2, 1):
	sld = ids[i]
	prs.part.drop_rel(sld.rId)
	ids.remove(sld)
thanks = ids[1]
ids.remove(thanks)
ids.append(thanks)
prs.save(sys.argv[2])
print('wrote', sys.argv[2], len(prs.slides), 'slides')
